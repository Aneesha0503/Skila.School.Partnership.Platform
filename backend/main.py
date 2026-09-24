import io
import csv
import json
import uuid
import os
import smtplib
from email.message import EmailMessage
from datetime import datetime, timezone
from typing import Optional, List, Dict, Any

from fastapi import FastAPI, HTTPException, Query, Response, Header, Depends
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from firebase_config import get_db, is_live_firebase
from models import SchoolModel, SchoolCreateUpdate, AgentNoteCreate, AgentNoteModel

app = FastAPI(title="Skila School Partnership Platform API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

from india_data import get_states_and_districts
from mistral_scraper import get_school_tier

db = get_db()

# ==========================================
# ROLE-BASED ACCESS CONTROL (RBAC) HELPERS
# ==========================================
def get_current_role(x_user_role: Optional[str] = Header(None)) -> str:
    """Extracts and validates user role from X-User-Role header. Defaults to 'admin'."""
    if not x_user_role:
        return "admin"
    role = x_user_role.strip().lower()
    if role not in ("admin", "agent"):
        return "admin"
    return role

def require_admin(role: str = Depends(get_current_role)) -> str:
    """Enforces that only users with 'admin' role can access the route."""
    if role != "admin":
        raise HTTPException(
            status_code=403,
            detail="Access Denied: Administrator role required to perform this action."
        )
    return role

@app.get("/api/auth/roles")
def get_roles_matrix(role: str = Depends(get_current_role)):
    """Returns available roles and full permissions matrix."""
    return {
        "active_role": role,
        "roles": {
            "admin": {
                "name": "Administrator",
                "badge": "Admin Level",
                "tag": "👑 Super Admin",
                "description": "Full unconstrained administrative authority across data scraping, AI models, deletion, and system configuration.",
                "permissions": {
                    "can_run_district_discovery": True,
                    "can_run_single_school_ai": True,
                    "can_add_school": True,
                    "can_edit_core_info": True,
                    "can_delete_school": True,
                    "can_clear_database": True,
                    "can_export_excel": True,
                    "can_update_crm": True
                }
            },
            "agent": {
                "name": "Field Agent",
                "badge": "Agent Level",
                "tag": "💼 Partnership Agent",
                "description": "Field intelligence access: browse directory, research institutional profiles, and update CRM lead milestones.",
                "permissions": {
                    "can_run_district_discovery": False,
                    "can_run_single_school_ai": False,
                    "can_add_school": False,
                    "can_edit_core_info": False,
                    "can_delete_school": False,
                    "can_clear_database": False,
                    "can_export_excel": True,
                    "can_update_crm": True
                }
            }
        }
    }


def get_all_schools_raw() -> List[Dict[str, Any]]:
    col = db.collection("schools")
    docs = col.stream()
    schools = []
    for d in docs:
        data = d.to_dict()
        if not data.get("id"):
            data["id"] = d.id
        if not data.get("tier"):
            data["tier"] = get_school_tier(data)
        schools.append(data)
    return schools

@app.get("/api/states-districts")
def api_states_districts():
    """Returns official Indian states and their districts."""
    return get_states_and_districts()

@app.get("/api/status")
def get_status():
    schools = get_all_schools_raw()
    return {
        "status": "healthy",
        "firebase_live": is_live_firebase(),
        "total_schools": len(schools),
        "database_type": "Firebase Firestore Live" if is_live_firebase() else "Local Firestore Document Store"
    }

@app.get("/api/hierarchy")
def get_hierarchy(
    state: Optional[str] = None,
    district: Optional[str] = None,
    revenue_division: Optional[str] = None,
    mandal: Optional[str] = None,
    local_body_name: Optional[str] = None
):
    schools = get_all_schools_raw()
    
    states = sorted(list(set(s["hierarchy"]["state"] for s in schools if s.get("hierarchy", {}).get("state"))))
    
    dist_filtered = [s for s in schools if not state or s.get("hierarchy", {}).get("state") == state]
    districts = sorted(list(set(s["hierarchy"]["district"] for s in dist_filtered if s.get("hierarchy", {}).get("district"))))
    
    rev_filtered = [s for s in dist_filtered if not district or s.get("hierarchy", {}).get("district") == district]
    revenue_divisions = sorted(list(set(s["hierarchy"]["revenue_division"] for s in rev_filtered if s.get("hierarchy", {}).get("revenue_division"))))
    
    mandal_filtered = [s for s in rev_filtered if not revenue_division or s.get("hierarchy", {}).get("revenue_division") == revenue_division]
    mandals = sorted(list(set(s["hierarchy"]["mandal"] for s in mandal_filtered if s.get("hierarchy", {}).get("mandal"))))
    
    lb_filtered = [s for s in mandal_filtered if not mandal or s.get("hierarchy", {}).get("mandal") == mandal]
    local_bodies = []
    seen_lbs = set()
    for s in lb_filtered:
        h = s.get("hierarchy", {})
        lb_name = h.get("local_body_name")
        lb_type = h.get("local_body_type", "Municipality")
        if lb_name and (lb_type, lb_name) not in seen_lbs:
            seen_lbs.add((lb_type, lb_name))
            local_bodies.append({"type": lb_type, "name": lb_name})
    local_bodies.sort(key=lambda x: x["name"])
    
    ward_filtered = [s for s in lb_filtered if not local_body_name or s.get("hierarchy", {}).get("local_body_name") == local_body_name]
    wards_villages = sorted(list(set(s["hierarchy"]["village_locality_ward"] for s in ward_filtered if s.get("hierarchy", {}).get("village_locality_ward"))))
    
    return {
        "states": states,
        "districts": districts,
        "revenue_divisions": revenue_divisions,
        "mandals": mandals,
        "local_bodies": local_bodies,
        "wards_villages": wards_villages
    }

@app.get("/api/schools")
def list_schools(
    state: Optional[str] = None,
    district: Optional[str] = None,
    revenue_division: Optional[str] = None,
    mandal: Optional[str] = None,
    local_body_name: Optional[str] = None,
    village_locality_ward: Optional[str] = None,
    search: Optional[str] = None,
    board: Optional[str] = None,
    lead_status: Optional[str] = None,
    skila_ai_potential: Optional[str] = None,
    technology_adoption_level: Optional[str] = None,
    tier: Optional[str] = None
):
    schools = get_all_schools_raw()
    filtered = []
    
    search_lower = search.strip().lower() if search else None
    
    for s in schools:
        h = s.get("hierarchy", {})
        info = s.get("info", {})
        tech = s.get("technology", {})
        sales = s.get("sales", {})
        t_info = s.get("tier", {})
        
        if state and h.get("state") != state:
            continue
        if district and h.get("district") != district:
            continue
        if revenue_division and h.get("revenue_division") != revenue_division:
            continue
        if mandal and h.get("mandal") != mandal:
            continue
        if local_body_name and h.get("local_body_name") != local_body_name:
            continue
        if village_locality_ward and h.get("village_locality_ward") != village_locality_ward:
            continue
            
        if tier and tier != "All" and t_info.get("tier") != tier:
            continue
        if board and board != "All" and board.lower() not in info.get("board", "").lower():
            continue
        if lead_status and lead_status != "All" and sales.get("lead_status") != lead_status:
            continue
        if skila_ai_potential and skila_ai_potential != "All" and sales.get("skila_ai_potential") != skila_ai_potential:
            continue
        if technology_adoption_level and technology_adoption_level != "All" and sales.get("technology_adoption_level") != technology_adoption_level:
            continue
            
        if search_lower:
            text_corpus = f"{info.get('school_name', '')} {info.get('udise_code', '')} {info.get('principal_name', '')} {info.get('correspondent_name', '')} {h.get('district', '')} {h.get('mandal', '')} {h.get('village_locality_ward', '')}".lower()
            if search_lower not in text_corpus:
                continue
                
        filtered.append(s)
        
    # Sort from High to Low: High Range first, then State Board by strength descending
    filtered.sort(key=lambda x: (
        x.get("tier", {}).get("rank", 99),
        -int(x.get("info", {}).get("student_strength") or 0)
    ))
    return filtered

@app.get("/api/schools/{school_id}")
def get_school(school_id: str):
    doc = db.collection("schools").document(school_id).get()
    if not doc.exists:
        raise HTTPException(status_code=404, detail="School not found")
    data = doc.to_dict()
    data["id"] = school_id
    return data

@app.post("/api/schools")
def create_school(payload: SchoolCreateUpdate, role: str = Depends(require_admin)):
    """Admin only: Create new school."""
    school_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    record = {
        "id": school_id,
        "hierarchy": payload.hierarchy.model_dump(exclude_unset=True) if payload.hierarchy else {},
        "info": payload.info.model_dump(exclude_unset=True) if payload.info else {},
        "technology": payload.technology.model_dump(exclude_unset=True) if payload.technology else {},
        "sales": payload.sales.model_dump(exclude_unset=True) if payload.sales else {},
        "created_at": now,
        "updated_at": now,
        "created_by_role": role
    }
    db.collection("schools").document(school_id).set(record)
    return record

@app.put("/api/schools/{school_id}")
def update_school(school_id: str, payload: SchoolCreateUpdate, role: str = Depends(get_current_role)):
    """
    Update school:
    - Admin: full update across hierarchy, school info, tech, and sales.
    - Agent: updates CRM sales milestones and remarks, preserving core hierarchy and UDISE.
    """
    doc_ref = db.collection("schools").document(school_id)
    doc = doc_ref.get()
    if not doc.exists:
        raise HTTPException(status_code=404, detail="School not found")
    
    existing = doc.to_dict()
    now = datetime.now(timezone.utc).isoformat()
    
    # Safely extract updates without overwriting existing nested data (e.g. sent_emails/sent_whatsapp)
    sales_update = payload.sales.model_dump(exclude_unset=True) if payload.sales else {}
    tech_update = payload.technology.model_dump(exclude_unset=True) if payload.technology else {}
    hierarchy_update = payload.hierarchy.model_dump(exclude_unset=True) if payload.hierarchy else {}
    info_update = payload.info.model_dump(exclude_unset=True) if payload.info else {}

    merged_sales = {**existing.get("sales", {}), **sales_update}
    merged_tech = {**existing.get("technology", {}), **tech_update}

    if role == "agent":
        # Agent level: allow updating sales milestones, notes, technology observations,
        # but preserve official administrative hierarchy and verified UDISE code
        record = {
            **existing,
            "id": school_id,
            "sales": merged_sales,
            "technology": merged_tech,
            "updated_at": now,
            "last_updated_by_role": "agent"
        }
    else:
        # Admin level: unconstrained update
        record = {
            **existing,
            "id": school_id,
            "hierarchy": {**existing.get("hierarchy", {}), **hierarchy_update} if hierarchy_update else existing.get("hierarchy", {}),
            "info": {**existing.get("info", {}), **info_update} if info_update else existing.get("info", {}),
            "technology": merged_tech,
            "sales": merged_sales,
            "created_at": existing.get("created_at", now),
            "updated_at": now,
            "last_updated_by_role": "admin"
        }
        
    doc_ref.set(record)
    return record

@app.post("/api/schools/{school_id}/agent-notes")
def add_agent_note(
    school_id: str,
    payload: AgentNoteCreate,
    role: str = Depends(get_current_role)
):
    """
    Agent / Admin logs a field note / update for a school.
    - Appends note to school.agent_notes
    - Generates an Admin alert notification in 'notifications' collection
    - Returns updated school and the note object
    """
    doc_ref = db.collection("schools").document(school_id)
    doc = doc_ref.get()
    if not doc.exists:
        raise HTTPException(status_code=404, detail="School not found")
        
    school = doc.to_dict()
    school["id"] = school_id
    now = datetime.now(timezone.utc).isoformat()
    note_id = str(uuid.uuid4())
    
    school_name = (
        school.get("info", {}).get("school_name") or 
        school.get("name") or 
        f"School {school_id[:8]}"
    )
    
    agent_display_name = payload.agent_name.strip() if payload.agent_name and payload.agent_name.strip() else ("Field Agent" if role == "agent" else "Admin")
    
    note_data = {
        "id": note_id,
        "school_id": school_id,
        "school_name": school_name,
        "agent_name": agent_display_name,
        "author_role": role,
        "category": payload.category or "School Visit",
        "urgency": payload.urgency or "Normal",
        "text": payload.text.strip(),
        "action_required": (payload.action_required or "").strip(),
        "timestamp": now,
        "admin_notified": True,
        "admin_read": False
    }
    
    # Append note to school record (newest first)
    existing_notes = school.get("agent_notes") or []
    updated_notes = [note_data] + existing_notes
    school["agent_notes"] = updated_notes
    school["updated_at"] = now
    school["last_updated_by_role"] = role
    
    doc_ref.set(school)
    
    # Save notification for Admin in 'notifications' collection
    notification_data = {
        "id": str(uuid.uuid4()),
        "note_id": note_id,
        "school_id": school_id,
        "school_name": school_name,
        "district": school.get("hierarchy", {}).get("district", ""),
        "state": school.get("hierarchy", {}).get("state", ""),
        "agent_name": note_data["agent_name"],
        "category": note_data["category"],
        "urgency": note_data["urgency"],
        "text_snippet": (note_data["text"][:140] + "...") if len(note_data["text"]) > 140 else note_data["text"],
        "full_text": note_data["text"],
        "timestamp": now,
        "is_read": False,
        "type": "agent_field_update"
    }
    
    try:
        db.collection("notifications").document(notification_data["id"]).set(notification_data)
    except Exception as e:
        print(f"Warning: Could not save notification to Firestore: {e}")
        
    return {
        "success": True,
        "note": note_data,
        "school": school,
        "notification": notification_data
    }

@app.get("/api/notifications")
def get_notifications(
    limit: int = 50,
    role: str = Depends(get_current_role)
):
    """
    Fetches recent agent update notifications for Admin.
    """
    try:
        col = db.collection("notifications")
        docs = col.stream()
        notifications = []
        for d in docs:
            item = d.to_dict()
            if not item.get("id"):
                item["id"] = d.id
            notifications.append(item)
            
        notifications.sort(key=lambda x: x.get("timestamp", ""), reverse=True)
        return notifications[:limit]
    except Exception as e:
        print(f"Error fetching notifications: {e}")
        return []

@app.put("/api/notifications/{notification_id}/read")
def mark_notification_read(notification_id: str):
    """Marks a single notification as read."""
    try:
        doc_ref = db.collection("notifications").document(notification_id)
        doc = doc_ref.get()
        if doc.exists:
            doc_ref.update({"is_read": True})
        return {"success": True, "id": notification_id}
    except Exception as e:
        return {"success": False, "error": str(e)}

@app.post("/api/notifications/mark-all-read")
def mark_all_notifications_read():
    """Marks all notifications as read."""
    try:
        col = db.collection("notifications")
        docs = col.stream()
        for d in docs:
            data = d.to_dict()
            if not data.get("is_read"):
                col.document(d.id).update({"is_read": True})
        return {"success": True}
    except Exception as e:
        return {"success": False, "error": str(e)}

@app.delete("/api/schools/clear-all")
def clear_all_schools(role: str = Depends(require_admin)):
    """Admin only: Clears all schools from the database."""
    col = db.collection("schools")
    docs = list(col.stream())
    count = len(docs)
    for d in docs:
        col.document(d.id).delete()
    return {"message": f"Successfully deleted {count} schools", "deleted_count": count}

@app.delete("/api/schools/{school_id}")
def delete_school(school_id: str, role: str = Depends(require_admin)):
    """Admin only: Delete single school."""
    doc_ref = db.collection("schools").document(school_id)
    doc = doc_ref.get()
    if not doc.exists:
        raise HTTPException(status_code=404, detail="School not found")
    doc_ref.delete()
    return {"message": "School successfully deleted", "id": school_id}

@app.get("/api/stats")
def get_stats(state: Optional[str] = None, district: Optional[str] = None):
    schools = get_all_schools_raw()
    if state:
        schools = [s for s in schools if s.get("hierarchy", {}).get("state", "").lower() == state.strip().lower()]
    if district:
        schools = [s for s in schools if s.get("hierarchy", {}).get("district", "").lower() == district.strip().lower()]
        
    total = len(schools)
    
    high_range = sum(1 for s in schools if s.get("tier", {}).get("tier") == "High Range")
    state_high = sum(1 for s in schools if s.get("tier", {}).get("tier") == "State Board - High Strength")
    state_mid = sum(1 for s in schools if s.get("tier", {}).get("tier") == "State Board - Mid Strength")
    state_low = sum(1 for s in schools if s.get("tier", {}).get("tier") == "State Board - Low Strength")
    
    total_students = sum(int(s.get("info", {}).get("student_strength") or 0) for s in schools)
    total_teachers = sum(int(s.get("info", {}).get("teacher_strength") or 0) for s in schools)
    
    high_potential = sum(1 for s in schools if s.get("sales", {}).get("skila_ai_potential") == "High")
    demos_done = sum(1 for s in schools if s.get("sales", {}).get("demo_done") == "Yes")
    proposals_shared = sum(1 for s in schools if s.get("sales", {}).get("proposal_shared") == "Yes")
    pilots_started = sum(1 for s in schools if s.get("sales", {}).get("pilot_started") == "Yes")
    
    return {
        "total_schools": total,
        "high_range": high_range,
        "state_high": state_high,
        "state_mid": state_mid,
        "state_low": state_low,
        "total_students": total_students,
        "total_teachers": total_teachers,
        "high_ai_potential": high_potential,
        "demos_done": demos_done,
        "proposals_shared": proposals_shared,
        "pilots_started": pilots_started
    }

def build_schools_excel(schools: List[Dict[str, Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schools Intelligence"
    
    # Enable grid lines
    ws.views.sheetView[0].showGridLines = True
    
    headers = [
        "#",
        "Tier Classification",
        "UDISE Code",
        "School Name",
        "Affiliated Board",
        "Student Strength",
        "Teacher Strength",
        "Student-Teacher Ratio",
        "State",
        "District",
        "Revenue Division",
        "Mandal",
        "Local Body Type",
        "Local Body Name",
        "Village / Locality / Ward",
        "Full Address",
        "Pincode",
        "Category",
        "Management Type",
        "School Type",
        "Classes From",
        "Classes To",
        "Principal Name",
        "Correspondent Name",
        "Mobile Number",
        "Official Email",
        "Website",
        "ERP Used",
        "ERP Vendor",
        "LMS Used",
        "LMS Vendor",
        "Coding Curriculum",
        "Coding Vendor",
        "Robotics Program",
        "Robotics Vendor",
        "AI Used",
        "AI Vendor",
        "STEM Program",
        "ATL Lab (NITI Aayog)",
        "Smart Classrooms",
        "Computer Labs",
        "Internet Connectivity",
        "Parent App",
        "School App",
        "Key Decision Maker",
        "Decision Maker Designation",
        "Decision Maker Contact",
        "Annual Fee Range",
        "Technology Adoption Level",
        "Skila AI Potential",
        "Lead Status",
        "Interest Level",
        "Demo Completed",
        "Proposal Shared",
        "Pilot Started",
        "Sales Owner",
        "Last Contact Date",
        "Next Follow-up Date",
        "Existing EdTech Partners",
        "Strategy Remarks",
        "Details Status"
    ]
    
    # Sort schools by tier and student strength
    tier_order = {
        "High Range": 1,
        "State Board - High Strength": 2,
        "State Board - Mid Strength": 3,
        "State Board - Low Strength": 4
    }
    
    def get_sort_key(s):
        t = s.get("tier", {}).get("tier", "Standard")
        rank = tier_order.get(t, 5)
        strength = int(s.get("info", {}).get("student_strength") or 0)
        return (rank, -strength)
        
    sorted_schools = sorted(schools, key=get_sort_key)
    
    # Header styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = cell_border
        
    row_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    
    for row_idx, s in enumerate(sorted_schools, start=2):
        h = s.get("hierarchy", {})
        info = s.get("info", {})
        tech = s.get("technology", {})
        sales = s.get("sales", {})
        tier = s.get("tier", {}).get("tier", "Standard")
        
        students = int(info.get("student_strength") or 0)
        teachers = int(info.get("teacher_strength") or 0)
        ratio = f"{(students / teachers):.1f}:1" if teachers > 0 else "N/A"
        
        partners = sales.get("existing_edtech_partners", "")
        if isinstance(partners, list):
            partners = ", ".join(partners)
            
        remarks = sales.get("remarks", "")
        if isinstance(remarks, dict):
            remarks = "; ".join(f"{k}: {v}" for k, v in remarks.items())
            
        details_status = "49 Fields Ready" if s.get("details_fetched") else "Pending Deep Run"
        
        row_values = [
            row_idx - 1,
            tier,
            str(info.get("udise_code") or ""),
            info.get("school_name", ""),
            info.get("board", ""),
            students,
            teachers,
            ratio,
            h.get("state", ""),
            h.get("district", ""),
            h.get("revenue_division", ""),
            h.get("mandal", ""),
            h.get("local_body_type", ""),
            h.get("local_body_name", ""),
            h.get("village_locality_ward", ""),
            info.get("full_address", ""),
            str(info.get("pincode") or ""),
            info.get("school_category", ""),
            info.get("management_type", ""),
            info.get("school_type", ""),
            info.get("classes_from", ""),
            info.get("classes_to", ""),
            info.get("principal_name", ""),
            info.get("correspondent_name", ""),
            str(info.get("mobile") or ""),
            info.get("email", ""),
            info.get("website", ""),
            tech.get("erp_used", ""),
            tech.get("erp_vendor", ""),
            tech.get("lms_used", ""),
            tech.get("lms_vendor", ""),
            tech.get("coding_used", ""),
            tech.get("coding_vendor", ""),
            tech.get("robotics_used", ""),
            tech.get("robotics_vendor", ""),
            tech.get("ai_used", ""),
            tech.get("ai_vendor", ""),
            tech.get("stem_program", ""),
            tech.get("atl_lab", ""),
            int(tech.get("smart_classroom_count") or 0),
            int(tech.get("computer_lab_count") or 0),
            tech.get("internet", ""),
            tech.get("parent_app", ""),
            tech.get("school_app", ""),
            sales.get("decision_maker", ""),
            sales.get("decision_maker_designation", ""),
            str(sales.get("decision_maker_contact") or ""),
            sales.get("annual_fee_range", ""),
            sales.get("technology_adoption_level", ""),
            sales.get("skila_ai_potential", ""),
            sales.get("lead_status", ""),
            sales.get("interest_level", ""),
            sales.get("demo_done", ""),
            sales.get("proposal_shared", ""),
            sales.get("pilot_started", ""),
            sales.get("sales_owner", ""),
            sales.get("last_contact_date", ""),
            sales.get("next_follow_up_date", ""),
            partners,
            str(remarks),
            details_status
        ]
        
        ws.append(row_values)
        ws.row_dimensions[row_idx].height = 20
        
        is_alt = (row_idx % 2 == 1)
        for col_idx in range(1, len(row_values) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = cell_border
            if is_alt:
                cell.fill = row_alt_fill
            if col_idx in [1, 6, 7, 40, 41]: # Numbers / counts
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in [2, 3, 5, 8, 17, 28, 30, 32, 34, 36, 38, 39, 43, 44, 50, 51, 52, 53, 54, 55, 61]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        adjusted_width = min(max(max_len + 4, 12), 40)
        ws.column_dimensions[col_letter].width = adjusted_width
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

@app.get("/api/export")
@app.get("/api/export/excel")
def export_excel(state: Optional[str] = None, district: Optional[str] = None):
    schools = get_all_schools_raw()
    if state:
        schools = [s for s in schools if s.get("hierarchy", {}).get("state", "").lower() == state.strip().lower()]
    if district:
        schools = [s for s in schools if s.get("hierarchy", {}).get("district", "").lower() == district.strip().lower()]
        
    excel_bytes = build_schools_excel(schools)
    
    filename = "skila_schools_intelligence.xlsx"
    if district:
        clean_dist = district.lower().replace(" ", "_")
        filename = f"skila_schools_{clean_dist}.xlsx"
        
    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition"
        }
    )

@app.get("/api/export/csv")
def export_csv(state: Optional[str] = None, district: Optional[str] = None):
    schools = get_all_schools_raw()
    if state:
        schools = [s for s in schools if s.get("hierarchy", {}).get("state", "").lower() == state.strip().lower()]
    if district:
        schools = [s for s in schools if s.get("hierarchy", {}).get("district", "").lower() == district.strip().lower()]
    output = io.StringIO()
    writer = csv.writer(output)
    
    headers = [
        "UDISE Code", "School Name", "State", "District", "Mandal", "Ward/Village",
        "Category", "Management", "Board", "Students", "Teachers", "Principal",
        "Mobile", "Email", "ERP Used", "LMS Used", "Coding", "Robotics", "AI Used",
        "Decision Maker", "Designation", "Fee Range", "AI Potential", "Lead Status", "Remarks"
    ]
    writer.writerow(headers)
    
    for s in schools:
        h = s.get("hierarchy", {})
        info = s.get("info", {})
        tech = s.get("technology", {})
        sales = s.get("sales", {})
        writer.writerow([
            info.get("udise_code", ""), info.get("school_name", ""),
            h.get("state", ""), h.get("district", ""), h.get("mandal", ""), h.get("village_locality_ward", ""),
            info.get("school_category", ""), info.get("management_type", ""), info.get("board", ""),
            info.get("student_strength", 0), info.get("teacher_strength", 0), info.get("principal_name", ""),
            info.get("mobile", ""), info.get("email", ""),
            tech.get("erp_used", ""), tech.get("lms_used", ""), tech.get("coding_used", ""), tech.get("robotics_used", ""), tech.get("ai_used", ""),
            sales.get("decision_maker", ""), sales.get("decision_maker_designation", ""),
            sales.get("annual_fee_range", ""), sales.get("skila_ai_potential", ""),
            sales.get("lead_status", ""), sales.get("remarks", "")
        ])
        
    output.seek(0)
    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=skila_schools.csv"}
    )

from mistral_scraper import (
    scrape_schools_ai, scrape_district_schools_ai, scrape_single_school_details_ai, 
    enrich_school_with_ai, generate_contextual_email_ai, generate_whatsapp_pitch_ai
)

class AIScrapeRequest(BaseModel):
    query: Optional[str] = None
    state: Optional[str] = "Telangana"
    district: Optional[str] = "Hyderabad"
    mandal: Optional[str] = None
    count: int = 4
    auto_save: bool = False

class AISaveScrapedRequest(BaseModel):
    schools: List[Dict[str, Any]]

@app.post("/api/ai/scrape")
def api_scrape_schools(req: AIScrapeRequest):
    """
    Triggers Skila AI Intelligence Engine to research, extract, and scrape
    schools matching the region or search prompt.
    """
    scraped = scrape_schools_ai(
        query=req.query,
        state=req.state,
        district=req.district,
        mandal=req.mandal,
        count=req.count
    )
    if req.auto_save:
        col = db.collection("schools")
        for s in scraped:
            col.document(s["id"]).set(s)
    return {"schools": scraped, "count": len(scraped), "auto_saved": req.auto_save}

@app.post("/api/ai/save")
def api_save_scraped_schools(req: AISaveScrapedRequest):
    """
    Saves selected scraped schools into Firebase Firestore.
    """
    col = db.collection("schools")
    saved_count = 0
    now = datetime.now(timezone.utc).isoformat()
    for s in req.schools:
        doc_id = s.get("id") or str(uuid.uuid4())
        s["id"] = doc_id
        if not s.get("created_at"):
            s["created_at"] = now
        s["updated_at"] = now
        col.document(doc_id).set(s)
        saved_count += 1
    return {"message": f"Successfully imported {saved_count} schools", "count": saved_count}

@app.post("/api/ai/enrich/{school_id}")
def api_enrich_school(school_id: str):
    """
    Uses Skila AI to evaluate an existing school and enrich its tech recommendations
    and sales strategy pitch.
    """
    doc_ref = db.collection("schools").document(school_id)
    doc = doc_ref.get()
    if not doc.exists:
        raise HTTPException(status_code=404, detail="School not found")
    school_data = doc.to_dict()
    school_data["id"] = school_id
    
    enriched = enrich_school_with_ai(school_data)
    
    if "remarks" in enriched and enriched["remarks"]:
        current_remarks = school_data.get("sales", {}).get("remarks", "")
        enrich_rem = enriched["remarks"]
        if isinstance(enrich_rem, dict):
            lines = ["[Skila AI Analysis]:"]
            for k, v in enrich_rem.items():
                title = k.replace("_", " ").title()
                if isinstance(v, list):
                    lines.append(f"  • {title}:")
                    for item in v:
                        lines.append(f"    - {item}")
                else:
                    lines.append(f"  • {title}: {v}")
            enrich_rem_str = "\n".join(lines)
        else:
            enrich_rem_str = f"[Skila AI Analysis]: {enrich_rem}"
        
        updated_remarks = f"{current_remarks}\n\n{enrich_rem_str}".strip()
        school_data["sales"]["remarks"] = updated_remarks
        if "skila_ai_potential" in enriched:
            school_data["sales"]["skila_ai_potential"] = enriched["skila_ai_potential"]
        if "technology_adoption_level" in enriched:
            school_data["sales"]["technology_adoption_level"] = enriched["technology_adoption_level"]
        school_data["updated_at"] = datetime.now(timezone.utc).isoformat()
        doc_ref.set(school_data)
        
    return {"school": school_data, "ai_insights": enriched}

class RunDistrictRequest(BaseModel):
    state: str
    district: str
    count: int = 25
    force_scrape: bool = False
    scrape_more: bool = False

@app.post("/api/run-district")
def api_run_district(req: RunDistrictRequest, role: str = Depends(require_admin)):
    """
    Step 1: Runs district school discovery for selected State and District.
    Restricted to Administrator role.
    Returns schools strictly sorted High to Low:
    1. High Range (International, Cambridge, CBSE, ICSE)
    2. State Board High Strength (1200+)
    3. State Board Mid Strength (500-1200)
    4. State Board Low Strength (<500)
    Supports continuous multi-mandal scaling (25, 50, 100+ schools) without duplicates.
    """
    state_clean = req.state.strip()
    district_clean = req.district.strip()
    col = db.collection("schools")
    
    existing = [
        s for s in get_all_schools_raw()
        if s.get("hierarchy", {}).get("state", "").lower() == state_clean.lower()
        and s.get("hierarchy", {}).get("district", "").lower() == district_clean.lower()
    ]
    
    # If force_scrape is requested, wipe existing schools for this district to start fresh
    if req.force_scrape:
        for s in existing:
            col.document(s["id"]).delete()
        existing = []

    # If scrape_more is requested or if fewer than req.count schools exist:
    if req.scrape_more or len(existing) < req.count:
        existing_names = [s.get("info", {}).get("school_name", "") for s in existing]
        needed_count = req.count if req.scrape_more else (req.count - len(existing))
        
        scraped = scrape_district_schools_ai(
            state=state_clean,
            district=district_clean,
            count=min(50, max(15, needed_count)),
            exclude_names=existing_names
        )
        
        for s in scraped:
            col.document(s["id"]).set(s)
            
        all_district_schools = [
            s for s in get_all_schools_raw()
            if s.get("hierarchy", {}).get("state", "").lower() == state_clean.lower()
            and s.get("hierarchy", {}).get("district", "").lower() == district_clean.lower()
        ]
        all_district_schools.sort(key=lambda x: (
            x.get("tier", {}).get("rank", 99),
            -int(x.get("info", {}).get("student_strength") or 0)
        ))
        
        return {
            "schools": all_district_schools,
            "count": len(all_district_schools),
            "source": "skila_ai",
            "state": state_clean,
            "district": district_clean,
            "scraped_new": len(scraped)
        }
        
    # Return existing schools sorted High to Low
    existing.sort(key=lambda x: (
        x.get("tier", {}).get("rank", 99),
        -int(x.get("info", {}).get("student_strength") or 0)
    ))
    return {
        "schools": existing,
        "count": len(existing),
        "source": "database",
        "state": state_clean,
        "district": district_clean,
        "scraped_new": 0
    }

@app.post("/api/schools/{school_id}/run-details")
def api_run_school_details(school_id: str, role: str = Depends(require_admin)):
    """
    Step 2: On-demand single school intelligence runner.
    Restricted to Administrator role.
    When user approves or selects a specific school from the district list,
    runs Skila AI specifically for that school to populate all 16 Info,
    17 Technology, and 16 Sales CRM fields.
    """
    doc_ref = db.collection("schools").document(school_id)
    doc = doc_ref.get()
    if not doc.exists:
        raise HTTPException(status_code=404, detail="School not found")
    school_data = doc.to_dict()
    school_data["id"] = school_id
    
    enriched = scrape_single_school_details_ai(school_data)
    doc_ref.set(enriched)
    return enriched

class SendEmailRequest(BaseModel):
    recipient_email: str
    recipient_name: Optional[str] = ""
    subject: str
    body: str
    sender_type: Optional[str] = "company"  # "company" | "personal"

@app.get("/api/email-config")
def api_get_email_config():
    """
    Returns the configured sender profiles (Company vs Personal)
    so the frontend dropdown can display active sender names and email addresses.
    """
    company_name = os.getenv("COMPANY_EMAIL_NAME", "Skila AI Partnerships")
    company_email = os.getenv("COMPANY_EMAIL_ADDRESS", os.getenv("SMTP_USER", "partnerships@skila.ai"))
    company_configured = bool(
        (os.getenv("COMPANY_SMTP_USER") and os.getenv("COMPANY_SMTP_PASSWORD")) or
        (os.getenv("SMTP_USER") and os.getenv("SMTP_PASSWORD"))
    )

    personal_name = os.getenv("PERSONAL_EMAIL_NAME", "Personal Representative")
    personal_email = os.getenv("PERSONAL_EMAIL_ADDRESS", os.getenv("PERSONAL_SMTP_USER", "personal@gmail.com"))
    personal_configured = bool(os.getenv("PERSONAL_SMTP_USER") and os.getenv("PERSONAL_SMTP_PASSWORD"))

    return {
        "company": {
            "id": "company",
            "name": company_name,
            "email": company_email,
            "is_configured": company_configured,
            "label": f"🏢 Company: {company_name} ({company_email})"
        },
        "personal": {
            "id": "personal",
            "name": personal_name,
            "email": personal_email,
            "is_configured": personal_configured,
            "label": f"👤 Personal: {personal_name} ({personal_email})"
        }
    }

@app.post("/api/schools/{school_id}/generate-email")
def api_generate_school_email(school_id: str):
    """
    Uses Skila AI to automatically generate a tailored, contextual
    partnership proposal email for this specific school.
    """
    doc_ref = db.collection("schools").document(school_id)
    doc = doc_ref.get()
    if not doc.exists:
        raise HTTPException(status_code=404, detail="School not found")
    school_data = doc.to_dict()
    school_data["id"] = school_id
    
    generated = generate_contextual_email_ai(school_data)
    return generated

@app.post("/api/schools/{school_id}/send-email")
def api_send_school_email(school_id: str, req: SendEmailRequest):
    """
    Dispatches a contextual partnership email to the school or principal.
    Supports routing through Company Account or Personal Account.
    """
    doc_ref = db.collection("schools").document(school_id)
    doc = doc_ref.get()
    if not doc.exists:
        raise HTTPException(status_code=404, detail="School not found")
    school_data = doc.to_dict()
    school_data["id"] = school_id

    recipient = req.recipient_email.strip()
    if not recipient:
        raise HTTPException(status_code=400, detail="Recipient email address is required")

    sender_type = (req.sender_type or "company").lower()

    if sender_type == "personal":
        smtp_host = os.getenv("PERSONAL_SMTP_HOST", os.getenv("SMTP_HOST", "smtp.gmail.com"))
        smtp_port = int(os.getenv("PERSONAL_SMTP_PORT", os.getenv("SMTP_PORT", 587)))
        smtp_user = os.getenv("PERSONAL_SMTP_USER")
        smtp_password = os.getenv("PERSONAL_SMTP_PASSWORD")
        smtp_from = os.getenv("PERSONAL_EMAIL_ADDRESS", smtp_user or "personal@gmail.com")
        sender_label = os.getenv("PERSONAL_EMAIL_NAME", "Personal Representative")
    else:
        smtp_host = os.getenv("COMPANY_SMTP_HOST", os.getenv("SMTP_HOST", "smtp.gmail.com"))
        smtp_port = int(os.getenv("COMPANY_SMTP_PORT", os.getenv("SMTP_PORT", 587)))
        smtp_user = os.getenv("COMPANY_SMTP_USER", os.getenv("SMTP_USER"))
        smtp_password = os.getenv("COMPANY_SMTP_PASSWORD", os.getenv("SMTP_PASSWORD"))
        smtp_from = os.getenv("COMPANY_EMAIL_ADDRESS", smtp_user or "partnerships@skila.ai")
        sender_label = os.getenv("COMPANY_EMAIL_NAME", "Skila AI Partnerships")

    if smtp_user:
        smtp_user = smtp_user.strip(' "\'')
    if smtp_password:
        smtp_password = smtp_password.replace(" ", "").strip(' "\'')
    if smtp_from:
        smtp_from = smtp_from.strip(' "\'')

    smtp_sent = False
    smtp_error = None

    if smtp_host and smtp_user and smtp_password:
        try:
            msg = EmailMessage()
            msg["Subject"] = req.subject
            msg["From"] = f"{sender_label} <{smtp_from}>"
            msg["To"] = recipient
            msg.set_content(req.body)

            with smtplib.SMTP(smtp_host, smtp_port, timeout=8) as server:
                server.starttls()
                server.login(smtp_user, smtp_password)
                server.send_message(msg)
            smtp_sent = True
        except Exception as e:
            print(f"[SMTP Send Error - {sender_type}] {e}")
            smtp_error = str(e)

    if not smtp_sent:
        err_lower = str(smtp_error or "").lower()
        if "timed out" in err_lower or "timeout" in err_lower:
            detail = (
                "Outbound SMTP connection timed out on port 587/465. "
                "Your Internet Provider (ISP) or local network blocks direct SMTP ports. "
                "Please click 'Send via Gmail Web' to dispatch directly with 1 click!"
            )
        else:
            detail = f"SMTP Delivery Failed: {smtp_error or 'Could not authenticate with mail server.'}"
        raise HTTPException(status_code=502, detail=detail)

    # Automatically record engagement in CRM pipeline
    sales = school_data.setdefault("sales", {})
    sales["lead_status"] = "Contacted"
    now_utc = datetime.now(timezone.utc)
    today_str = now_utc.strftime("%Y-%m-%d")
    timestamp_str = now_utc.strftime("%Y-%m-%d %H:%M UTC")
    sales["last_contact_date"] = today_str

    current_remarks = sales.get("remarks", "")
    log_entry = f"[Email Sent via SMTP from {sender_label} ({smtp_from}) to {recipient} on {timestamp_str}]: {req.subject}"
    sales["remarks"] = f"{current_remarks}\n\n{log_entry}".strip()

    # Track in sent emails log array
    sent_list = sales.setdefault("sent_emails", [])
    sent_list.append({
        "timestamp": timestamp_str,
        "sender_type": sender_type,
        "sender_email": smtp_from,
        "sender_name": sender_label,
        "recipient_email": recipient,
        "recipient_name": req.recipient_name,
        "subject": req.subject,
        "smtp_sent": True
    })

    school_data["updated_at"] = now_utc.isoformat()
    doc_ref.set(school_data)

    return {
        "status": "success",
        "sender_type": sender_type,
        "sender_email": smtp_from,
        "sender_name": sender_label,
        "smtp_sent": True,
        "recipient": recipient,
        "message": f"Email successfully delivered to {recipient} from {sender_label} ({smtp_from})!",
        "school": school_data
    }

@app.post("/api/schools/{school_id}/log-email")
def api_log_school_email(school_id: str, req: SendEmailRequest):
    """
    Logs an email dispatched via Gmail Web or local client
    into the CRM pipeline and updates status to Contacted.
    """
    doc_ref = db.collection("schools").document(school_id)
    doc = doc_ref.get()
    if not doc.exists:
        raise HTTPException(status_code=404, detail="School not found")
    school_data = doc.to_dict()
    school_data["id"] = school_id

    recipient = req.recipient_email.strip()
    sender_type = (req.sender_type or "personal").lower()
    
    sender_label = os.getenv("PERSONAL_EMAIL_NAME", "Skila AI") if sender_type == "personal" else os.getenv("COMPANY_EMAIL_NAME", "Skila AI Partnerships")
    sender_email = os.getenv("PERSONAL_EMAIL_ADDRESS", "skila.udaymerugu@gmail.com") if sender_type == "personal" else os.getenv("COMPANY_EMAIL_ADDRESS", "partnerships@skila.ai")

    sales = school_data.setdefault("sales", {})
    sales["lead_status"] = "Contacted"
    now_utc = datetime.now(timezone.utc)
    today_str = now_utc.strftime("%Y-%m-%d")
    timestamp_str = now_utc.strftime("%Y-%m-%d %H:%M UTC")
    sales["last_contact_date"] = today_str

    current_remarks = sales.get("remarks", "")
    log_entry = f"[Email Dispatched via Gmail Web from {sender_label} ({sender_email}) to {recipient} on {timestamp_str}]: {req.subject}"
    sales["remarks"] = f"{current_remarks}\n\n{log_entry}".strip()

    sent_list = sales.setdefault("sent_emails", [])
    sent_list.append({
        "timestamp": timestamp_str,
        "sender_type": sender_type,
        "sender_email": sender_email,
        "sender_name": sender_label,
        "recipient_email": recipient,
        "recipient_name": req.recipient_name,
        "subject": req.subject,
        "method": "gmail_web"
    })

    school_data["updated_at"] = now_utc.isoformat()
    doc_ref.set(school_data)

    return {
        "status": "success",
        "message": f"Recorded outreach to {recipient} in CRM as Contacted!",
        "school": school_data
    }

# ==============================================================================
# WHATSAPP OUTREACH INTEGRATION (MSG91 + WHATSAPP WEB)
# ==============================================================================

class SendWhatsAppRequest(BaseModel):
    recipient_phone: str
    recipient_name: Optional[str] = ""
    message: str
    var1: Optional[str] = ""
    var2: Optional[str] = ""
    var3: Optional[str] = ""

@app.get("/api/whatsapp-config")
def api_get_whatsapp_config():
    """
    Returns WhatsApp sender status and integrated phone number.
    """
    authkey = os.getenv("MSG91_AUTHKEY")
    integrated_number = os.getenv("MSG91_INTEGRATED_NUMBER", "919390875225")
    template_name = os.getenv("MSG91_WHATSAPP_TEMPLATE", "skila_school_partnership")
    is_configured = bool(authkey and integrated_number)
    return {
        "is_configured": is_configured,
        "integrated_number": integrated_number,
        "template_name": template_name,
        "sender_name": "Skila AI",
        "provider": "MSG91"
    }

@app.post("/api/schools/{school_id}/generate-whatsapp")
def api_generate_school_whatsapp(school_id: str):
    """
    Uses Mistral AI to draft a tailored, high-conversion WhatsApp pitch
    for this specific school leadership.
    """
    doc_ref = db.collection("schools").document(school_id)
    doc = doc_ref.get()
    if not doc.exists:
        raise HTTPException(status_code=404, detail="School not found")
    school_data = doc.to_dict()
    school_data["id"] = school_id
    
    return generate_whatsapp_pitch_ai(school_data)

@app.post("/api/schools/{school_id}/send-whatsapp")
def api_send_school_whatsapp(school_id: str, req: SendWhatsAppRequest):
    """
    Dispatches automated WhatsApp outreach to the school leadership via MSG91 WhatsApp API.
    """
    doc_ref = db.collection("schools").document(school_id)
    doc = doc_ref.get()
    if not doc.exists:
        raise HTTPException(status_code=404, detail="School not found")
    school_data = doc.to_dict()
    school_data["id"] = school_id

    authkey = os.getenv("MSG91_AUTHKEY")
    integrated_number = os.getenv("MSG91_INTEGRATED_NUMBER", "919390875225")
    template_name = os.getenv("MSG91_WHATSAPP_TEMPLATE", "school_notice_general")
    namespace = os.getenv("MSG91_WHATSAPP_NAMESPACE", "da3632e0_cc65_4531_af19_8f0e0c271485")

    raw_phone = "".join(filter(str.isdigit, req.recipient_phone))
    if raw_phone.startswith("0"):
        raw_phone = raw_phone[1:]
    if len(raw_phone) == 10:
        raw_phone = f"91{raw_phone}"

    if not raw_phone:
        raise HTTPException(status_code=400, detail="Valid recipient phone number is required")

    if not authkey:
        raise HTTPException(status_code=400, detail="MSG91 AuthKey is not configured in backend/.env")

    # Call MSG91 WhatsApp Outbound API
    url = "https://control.msg91.com/api/v5/whatsapp/whatsapp-outbound-message/"
    payload = {
        "integrated_number": integrated_number,
        "content_type": "template",
        "payload": {
            "messaging_product": "whatsapp",
            "to": raw_phone,
            "type": "template",
            "template": {
                "name": template_name,
                "namespace": namespace,
                "language": {
                    "code": "en",
                    "policy": "deterministic"
                },
                "to_and_components": [
                    {
                        "to": [raw_phone],
                        "components": {
                            "body_1": {"type": "text", "value": req.var1 or req.recipient_name or "Principal"},
                            "body_2": {"type": "text", "value": req.var2 or school_data.get("info", {}).get("school_name", "School")},
                            "body_3": {"type": "text", "value": req.var3 or school_data.get("info", {}).get("board", "CBSE")}
                        }
                    }
                ]
            }
        }
    }

    headers = {
        "authkey": authkey,
        "Content-Type": "application/json",
        "Accept": "application/json"
    }

    try:
        import urllib.request
        data_bytes = json.dumps(payload).encode("utf-8")
        request = urllib.request.Request(url, data=data_bytes, headers=headers, method="POST")
        with urllib.request.urlopen(request, timeout=12) as response:
            res_body = response.read().decode("utf-8")
            res_json = json.loads(res_body) if res_body else {}
    except urllib.error.HTTPError as e:
        err_msg = e.read().decode("utf-8")
        if "418" in err_msg:
            raise HTTPException(
                status_code=403,
                detail="MSG91 IP Restriction (Code 418): Please disable 'API Security / IP Whitelist' on your MSG91 Authkey or whitelist IP 4.240.107.236 in your MSG91 panel."
            )
        raise HTTPException(status_code=e.code, detail=f"MSG91 API Error: {err_msg}")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Network error communicating with MSG91: {str(e)}")

    # Update CRM pipeline
    sales = school_data.setdefault("sales", {})
    sales["lead_status"] = "Contacted"
    now_utc = datetime.now(timezone.utc)
    today_str = now_utc.strftime("%Y-%m-%d")
    timestamp_str = now_utc.strftime("%Y-%m-%d %H:%M UTC")
    sales["last_contact_date"] = today_str

    current_remarks = sales.get("remarks", "")
    log_entry = f"[WhatsApp Sent via MSG91 from {integrated_number} to +{raw_phone} on {timestamp_str}]: Template {template_name}"
    sales["remarks"] = f"{current_remarks}\n\n{log_entry}".strip()

    sent_list = sales.setdefault("sent_whatsapp", [])
    sent_list.append({
        "timestamp": timestamp_str,
        "sender_phone": integrated_number,
        "recipient_phone": raw_phone,
        "recipient_name": req.recipient_name,
        "template": template_name,
        "status": "dispatched"
    })

    school_data["updated_at"] = now_utc.isoformat()
    doc_ref.set(school_data)

    return {
        "status": "success",
        "message": f"WhatsApp pitch successfully dispatched to +{raw_phone} via MSG91!",
        "response": res_json,
        "school": school_data
    }

@app.post("/api/schools/{school_id}/log-whatsapp")
def api_log_school_whatsapp(school_id: str, req: SendWhatsAppRequest):
    """
    Logs WhatsApp outreach dispatched via WhatsApp Web / WhatsApp Desktop app
    into the CRM pipeline.
    """
    doc_ref = db.collection("schools").document(school_id)
    doc = doc_ref.get()
    if not doc.exists:
        raise HTTPException(status_code=404, detail="School not found")
    school_data = doc.to_dict()
    school_data["id"] = school_id

    raw_phone = "".join(filter(str.isdigit, req.recipient_phone))
    sales = school_data.setdefault("sales", {})
    sales["lead_status"] = "Contacted"
    now_utc = datetime.now(timezone.utc)
    today_str = now_utc.strftime("%Y-%m-%d")
    timestamp_str = now_utc.strftime("%Y-%m-%d %H:%M UTC")
    sales["last_contact_date"] = today_str

    current_remarks = sales.get("remarks", "")
    log_entry = f"[WhatsApp Pitch Dispatched via WhatsApp Web to +{raw_phone} on {timestamp_str}]"
    sales["remarks"] = f"{current_remarks}\n\n{log_entry}".strip()

    sent_list = sales.setdefault("sent_whatsapp", [])
    sent_list.append({
        "timestamp": timestamp_str,
        "recipient_phone": raw_phone,
        "recipient_name": req.recipient_name,
        "method": "whatsapp_web"
    })

    school_data["updated_at"] = now_utc.isoformat()
    doc_ref.set(school_data)

    return {
        "status": "success",
        "message": f"Recorded WhatsApp outreach to +{raw_phone} in CRM as Contacted!",
        "school": school_data
    }


